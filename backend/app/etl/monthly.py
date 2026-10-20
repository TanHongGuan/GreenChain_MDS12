import csv
import io
import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation

from backend.app.etl.errors import ETLValidationError
from backend.app.etl.readers import ParsedSheet, read_csv_rows, read_xlsx_rows

CANONICAL_COLUMNS = ("reporting_period", "electricity_kwh", "water_m3", "co2_tco2e")
ALTERNATIVE_COLUMNS = ("month", "electricity_mwh", "water_litres", "carbon_kgco2e")
LEGACY_COLUMNS = {"metric_name", "value", "unit"}
MONTH_NAMES = {name: index for index, name in enumerate(
    ("jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"), 1
)}


@dataclass(frozen=True)
class MonthlyRow:
    reporting_period: str
    electricity_kwh: Decimal
    water_m3: Decimal
    co2_tco2e: Decimal


def read_input(raw_bytes: bytes, extension: str) -> ParsedSheet:
    if extension == ".csv":
        return read_csv_rows(raw_bytes)
    if extension != ".xlsx":
        raise ETLValidationError("Only CSV and XLSX files are supported.")
    # Monthly workbooks must identify the data sheet explicitly. The legacy
    # single-sheet convention remains available for existing uploads.
    from openpyxl import load_workbook

    from backend.app.etl.errors import ETLParseError

    try:
        workbook = load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise ETLParseError("XLSX file could not be read.") from exc
    try:
        if "Monthly_Data" in workbook.sheetnames:
            return read_xlsx_rows(raw_bytes, worksheet_name="Monthly_Data")
        headers, rows = read_xlsx_rows(raw_bytes)
        if LEGACY_COLUMNS.issubset(headers):
            return headers, rows
        raise ETLValidationError(
            "Monthly XLSX files need a worksheet named Monthly_Data with the documented monthly columns. "
            "Legacy workbooks need metric_name, value, and unit on the first sheet."
        )
    finally:
        workbook.close()


def classify_and_clean(headers: list[str], raw_rows: list[dict]) -> tuple[str, list[MonthlyRow] | None]:
    if len(headers) != len(set(headers)):
        raise ETLValidationError("Duplicate column names are not allowed.")
    names = set(headers)
    if LEGACY_COLUMNS.issubset(names):
        return "LEGACY", None
    if names & LEGACY_COLUMNS and not names & (set(CANONICAL_COLUMNS) | set(ALTERNATIVE_COLUMNS)):
        return "LEGACY", None
    if names == set(CANONICAL_COLUMNS):
        format_name, columns = "MONTHLY_STANDARD", CANONICAL_COLUMNS
    elif names == set(ALTERNATIVE_COLUMNS):
        format_name, columns = "MONTHLY_ALTERNATIVE", ALTERNATIVE_COLUMNS
    else:
        if names & (set(CANONICAL_COLUMNS) | set(ALTERNATIVE_COLUMNS)):
            missing = sorted((set(CANONICAL_COLUMNS) if "reporting_period" in names else set(ALTERNATIVE_COLUMNS)) - names)
            extra = sorted(names - set(CANONICAL_COLUMNS) - set(ALTERNATIVE_COLUMNS))
            raise ETLValidationError(f"Monthly columns do not match a supported template. Missing: {missing}; unexpected: {extra}.")
        raise ETLValidationError(
            "Unsupported columns. Use reporting_period,electricity_kwh,water_m3,co2_tco2e "
            "or Month,Electricity (MWh),Water (Litres),Carbon (kgCO2e)."
        )
    if not raw_rows:
        raise ETLValidationError("File does not contain any data rows.")
    rows: list[MonthlyRow] = []
    seen: set[str] = set()
    for line, source in enumerate(raw_rows, 2):
        period = _parse_month(source.get(columns[0]), line, columns[0], format_name)
        if period in seen:
            raise ETLValidationError(f"Row {line}: duplicate reporting_period {period}.")
        seen.add(period)
        values = [_number(source.get(name), line, name) for name in columns[1:]]
        if format_name == "MONTHLY_ALTERNATIVE":
            values = [values[0] * 1000, values[1] / 1000, values[2] / 1000]
        if any(value.as_tuple().exponent < -8 or value >= Decimal("1000000000000000") for value in values):
            raise ETLValidationError(f"Row {line}: converted value exceeds supported precision or range.")
        rows.append(MonthlyRow(period, *values))
    return format_name, sorted(rows, key=lambda row: row.reporting_period)


def _parse_month(raw: object, line: int, column: str, format_name: str) -> str:
    value = str(raw or "").strip()
    if format_name == "MONTHLY_STANDARD":
        if re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", value):
            return value
    else:
        match = re.fullmatch(r"([A-Za-z]{3})-(\d{4})", value)
        if match and match.group(1).lower() in MONTH_NAMES:
            return f"{match.group(2)}-{MONTH_NAMES[match.group(1).lower()]:02d}"
    raise ETLValidationError(f"Row {line}, {column}: invalid month '{value}'. Use YYYY-MM for standard or Jan-2026 for alternative input.")


def _number(raw: object, line: int, column: str) -> Decimal:
    value = str(raw if raw is not None else "").strip()
    if not value:
        raise ETLValidationError(f"Row {line}, {column}: a measurement is required.")
    try:
        number = Decimal(value)
    except InvalidOperation as exc:
        raise ETLValidationError(f"Row {line}, {column}: '{value}' is not numeric.") from exc
    if not number.is_finite() or number < 0:
        raise ETLValidationError(f"Row {line}, {column}: use a finite, non-negative number.")
    if number.as_tuple().exponent < -8 or number >= Decimal("1000000000000000"):
        raise ETLValidationError(f"Row {line}, {column}: number exceeds supported precision or range.")
    return number


def require_submission_quarter(rows: list[MonthlyRow], submission_period: str) -> None:
    if not re.fullmatch(r"\d{4}-Q[1-4]", submission_period):
        raise ETLValidationError("Monthly uploads require a submission reporting period such as 2026-Q1.")
    for row in rows:
        year, month = (int(part) for part in row.reporting_period.split("-"))
        expected = f"{year}-Q{(month - 1) // 3 + 1}"
        if expected != submission_period:
            raise ETLValidationError(f"Month {row.reporting_period} is outside submission period {submission_period}.")


def build_monthly_csv(rows: list[MonthlyRow]) -> bytes:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(CANONICAL_COLUMNS)
    for row in rows:
        writer.writerow((row.reporting_period, row.electricity_kwh, row.water_m3, row.co2_tco2e))
    return output.getvalue().encode("utf-8")
