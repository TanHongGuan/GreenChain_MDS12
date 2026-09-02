import csv
import io

from backend.app.etl.errors import ETLParseError
from backend.app.etl.schema import normalize_header

RawRow = dict[str, object]
ParsedSheet = tuple[list[str], list[RawRow]]


def read_csv_rows(raw_bytes: bytes) -> ParsedSheet:
    try:
        text = raw_bytes.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ETLParseError("CSV file is not valid UTF-8 text.") from exc

    reader = csv.reader(io.StringIO(text))
    try:
        header_row = next(reader)
    except StopIteration as exc:
        raise ETLParseError("CSV file is empty.") from exc

    if not header_row or all(not cell.strip() for cell in header_row):
        raise ETLParseError("CSV file has no header row.")

    headers = [normalize_header(cell) for cell in header_row]

    rows: list[RawRow] = []
    for line_number, row in enumerate(reader, start=2):
        if not row or all(not cell.strip() for cell in row):
            continue
        if len(row) != len(headers):
            raise ETLParseError(f"CSV row {line_number} does not match the header column count.")
        rows.append(dict(zip(headers, row, strict=True)))

    return headers, rows


def read_xlsx_rows(raw_bytes: bytes) -> ParsedSheet:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency is required by requirements.txt
        raise ETLParseError("XLSX support is unavailable.") from exc

    try:
        workbook = load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise ETLParseError("XLSX file could not be read.") from exc

    try:
        worksheet = workbook.active
        if worksheet is None:
            raise ETLParseError("XLSX file has no worksheets.")

        row_iter = worksheet.iter_rows(values_only=True)
        try:
            header_row = next(row_iter)
        except StopIteration as exc:
            raise ETLParseError("XLSX file is empty.") from exc

        if header_row is None or all(cell is None for cell in header_row):
            raise ETLParseError("XLSX file has no header row.")

        headers = [normalize_header(str(cell)) if cell is not None else "" for cell in header_row]

        rows: list[RawRow] = []
        for row in row_iter:
            if row is None or all(cell is None for cell in row):
                continue
            rows.append(dict(zip(headers, row, strict=False)))

        return headers, rows
    finally:
        workbook.close()
