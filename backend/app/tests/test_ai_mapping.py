import csv
from io import BytesIO
from decimal import Decimal
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch

import pytest
from openpyxl import Workbook

from backend.app.etl.ai_mapping import (
    ColumnMapping, MappingSuggestion, MappingUnavailable, inspect_sheets,
    inspection_payload, process_confirmed_mapping, suggest_mapping, validate_references,
)
from backend.app.etl.errors import ETLValidationError
from backend.app.models.monthly_record import MonthlyRecord
from backend.app.models.organisation import Organisation
from backend.app.models.project import Project
from backend.app.models.submission import Submission
from backend.app.submissions.service import process_submission
from backend.app.tests.test_submissions_service import db, settings, storage, upload, uploader_id


UNKNOWN = b"Billing Month,TNB Usage (MWh),Water Bill (L),GHG Emissions (kgCO2e)\nJan 2026,30,850000,14200\nFeb 2026,32,910000,15100\nMar 2026,28,780000,13600\n"


def suggestion(sheet="CSV"):
    return MappingSuggestion(mappings=[
        ColumnMapping(field="reporting_period", worksheet=sheet, header_row=1, header="Billing Month", source_unit="month"),
        ColumnMapping(field="electricity_kwh", worksheet=sheet, header_row=1, header="TNB Usage (MWh)", source_unit="MWh"),
        ColumnMapping(field="water_m3", worksheet=sheet, header_row=1, header="Water Bill (L)", source_unit="L"),
        ColumnMapping(field="co2_tco2e", worksheet=sheet, header_row=1, header="GHG Emissions (kgCO2e)", source_unit="kgCO2e"),
    ])


def test_confirmed_ai_mapping_produces_exact_canonical_measurements():
    rows = process_confirmed_mapping(UNKNOWN, ".csv", suggestion())
    assert [(row.reporting_period, str(row.electricity_kwh), str(row.water_m3), str(row.co2_tco2e)) for row in rows] == [
        ("2026-01", "30000", "850", "14.2"), ("2026-02", "32000", "910", "15.1"),
        ("2026-03", "28000", "780", "13.6"),
    ]


def test_wrong_ai_reference_is_flagged_for_user_correction():
    proposed = suggestion()
    proposed.mappings[1].header = "Invented column"
    checked = validate_references(proposed, inspect_sheets(UNKNOWN, ".csv", sample_only=True))
    assert checked.mappings[1].missing
    with pytest.raises(ETLValidationError):
        process_confirmed_mapping(UNKNOWN, ".csv", checked)


def test_missing_or_ambiguous_mapping_cannot_be_confirmed():
    proposed = suggestion()
    proposed.mappings[2].missing = True
    with pytest.raises(ETLValidationError, match="unambiguous"):
        process_confirmed_mapping(UNKNOWN, ".csv", proposed)


def test_unsupported_unit_is_rejected():
    proposed = suggestion()
    proposed.mappings[2].source_unit = "gallons"
    with pytest.raises(ETLValidationError, match="Unsupported unit"):
        process_confirmed_mapping(UNKNOWN, ".csv", proposed)


def test_multiple_worksheets_are_inspected_but_cross_sheet_mapping_is_rejected():
    book = Workbook()
    book.active.title = "Notes"
    data = book.create_sheet("Meter_Data")
    for row in [line.split(",") for line in UNKNOWN.decode().splitlines()]:
        data.append(row)
    buffer = BytesIO()
    book.save(buffer)
    sheets = inspect_sheets(buffer.getvalue(), ".xlsx", sample_only=True)
    assert [sheet.name for sheet in sheets] == ["Notes", "Meter_Data"]
    proposed = suggestion("Meter_Data")
    rows = process_confirmed_mapping(buffer.getvalue(), ".xlsx", proposed)
    assert rows[0].electricity_kwh == 30000
    proposed.mappings[1].worksheet = "Notes"
    with pytest.raises(ETLValidationError):
        process_confirmed_mapping(buffer.getvalue(), ".xlsx", proposed)


def test_only_header_columns_are_selectable_and_metadata_is_excluded():
    book = Workbook()
    notes = book.active
    notes.title = "00_Read_Me"
    notes.append(["Key", "Value"])
    notes.append(["Reporting period", "2026-Q1"])
    readings = book.create_sheet("Readings")
    readings.append(["Instructions for uploader"])
    readings.append(["Month", "Electricity (kWh)", "Water (m3)", "CO2 (tCO2e)"])
    readings.append(["2026-01", 10, 2, 1])
    buffer = BytesIO()
    book.save(buffer)
    payload = inspection_payload(inspect_sheets(buffer.getvalue(), ".xlsx", sample_only=True))
    assert payload[0]["candidates"] == []
    candidate = payload[1]["candidates"][0]
    assert candidate["header_row"] == 2
    assert candidate["columns"][1] == {"header": "Electricity (kWh)", "column_ref": "B"}
    assert "2026-01" not in [column["header"] for column in candidate["columns"]]
    assert candidate["sample_rows"][0][0] == "2026-01"


def test_monthly_metrics_align_across_separate_worksheets():
    book = Workbook()
    for name, heading, values in [
        ("Period", "Month", ["2026-01", "2026-02"]),
        ("Electric", "Grid (kWh)", [10, 20]),
        ("Water", "Usage (m3)", [3, 4]),
        ("Carbon", "Emissions (tCO2e)", [1, 2]),
    ]:
        sheet = book.active if name == "Period" else book.create_sheet(name)
        sheet.title = name
        sheet.append(["Month", heading, "Reference"] if name != "Period" else ["Month", "Other", "Value"])
        for index, value in enumerate(values):
            sheet.append([f"2026-0{index + 1}", value, "source"] if name != "Period" else [value, 0, 0])
    buffer = BytesIO()
    book.save(buffer)
    proposed = MappingSuggestion(mappings=[
        ColumnMapping(field="reporting_period", worksheet="Period", header="Month", source_unit="month"),
        ColumnMapping(field="electricity_kwh", worksheet="Electric", header="Grid (kWh)", period_column_ref="A", source_unit="kWh"),
        ColumnMapping(field="water_m3", worksheet="Water", header="Usage (m3)", period_column_ref="A", source_unit="m3"),
        ColumnMapping(field="co2_tco2e", worksheet="Carbon", header="Emissions (tCO2e)", period_column_ref="A", source_unit="tCO2e"),
    ])
    assert process_confirmed_mapping(buffer.getvalue(), ".xlsx", proposed)[1].electricity_kwh == 20
    book["Water"]["A3"] = "2026-03"
    buffer = BytesIO()
    book.save(buffer)
    with pytest.raises(ETLValidationError, match="do not align"):
        process_confirmed_mapping(buffer.getvalue(), ".xlsx", proposed)


@pytest.mark.parametrize("filename", ["D01_Bayview_2026_Q1_VALID.xlsx", "D02_Bayview_2026_Q2_VALID.xlsx", "D05_Bayview_2026_Q1_INVALID.xlsx"])
def test_actual_bayview_daily_workbooks_are_detected_but_not_misaggregated(filename):
    # The workbooks live in the repository so this runs for everyone, including CI.
    # They were previously read from one developer's Desktop, so the test skipped
    # silently everywhere else and never actually guarded anything.
    path = Path(__file__).resolve().parents[3] / "test_datasets" / "workbooks" / filename
    if not path.exists():
        pytest.skip(f"Workbook fixture missing: {path}")
    content = path.read_bytes()
    payload = inspection_payload(inspect_sheets(content, ".xlsx", sample_only=True))
    assert [sheet["worksheet"] for sheet in payload if sheet["candidates"]] == [
        "Electricity_Log", "Water_Readings", "Carbon_Register",
    ]
    assert payload[0]["candidates"][0]["columns"][1]["header"] == "Main Incoming Grid (kWh)"
    assert payload[3]["candidates"] == []
    proposed = MappingSuggestion(mappings=[
        ColumnMapping(field="reporting_period", worksheet="Carbon_Register", header="Month", source_unit="month"),
        ColumnMapping(field="electricity_kwh", worksheet="Electricity_Log", header="Main Incoming Grid (kWh)", period_column_ref="A", source_unit="kWh"),
        ColumnMapping(field="water_m3", worksheet="Water_Readings", header="Water Meter - cumulative (m3)", period_column_ref="A", source_unit="m3"),
        ColumnMapping(field="co2_tco2e", worksheet="Carbon_Register", header="Reported Scope 2 Electricity Emissions (tCO2e)", period_column_ref="A", source_unit="tCO2e"),
    ])
    # Without a confirmed aggregation rule the daily rows are still refused, and the
    # message now names the actual choice rather than stating a flat limitation.
    with pytest.raises(ETLValidationError, match="daily consumption to be summed, or cumulative meter readings"):
        process_confirmed_mapping(content, ".xlsx", proposed)


def test_bayview_processes_once_the_aggregation_rules_are_confirmed():
    """The D01 workbook end to end, using the rules its own Meter_And_ETL_Rules sheet states.

    Electricity is daily consumption on the inclusive main meter, water is a
    cumulative meter needing differencing from the stated opening reading of
    23410, and emissions are already monthly.
    """
    path = Path(__file__).resolve().parents[3] / "test_datasets" / "workbooks" / "D01_Bayview_2026_Q1_VALID.xlsx"
    if not path.exists():
        pytest.skip(f"Workbook fixture missing: {path}")
    content = path.read_bytes()
    confirmed = MappingSuggestion(mappings=[
        ColumnMapping(field="reporting_period", worksheet="Carbon_Register", header="Month", source_unit="month"),
        ColumnMapping(field="electricity_kwh", worksheet="Electricity_Log", header="Main Incoming Grid (kWh)",
                      period_column_ref="A", source_unit="kWh", aggregation="SUM_DAILY"),
        ColumnMapping(field="water_m3", worksheet="Water_Readings", header="Water Meter - cumulative (m3)",
                      period_column_ref="A", source_unit="m3", aggregation="DIFFERENCE_CUMULATIVE",
                      opening_reading=Decimal("23410")),
        ColumnMapping(field="co2_tco2e", worksheet="Carbon_Register",
                      header="Reported Scope 2 Electricity Emissions (tCO2e)", period_column_ref="A",
                      source_unit="tCO2e", additional_column_refs=["C"]),
    ])
    rows = process_confirmed_mapping(content, ".xlsx", confirmed)
    assert [row.reporting_period for row in rows] == ["2026-01", "2026-02", "2026-03"]

    # The strongest available check: the dataset pack ships the quarterly totals it
    # expects, written independently of this implementation.
    expected_csv = (Path(__file__).resolve().parents[3] / "test_datasets" / "workbooks"
                    / "expected" / "D01_expected_quarterly.csv")
    header, values = expected_csv.read_text(encoding="utf-8-sig").strip().splitlines()
    _, _, electricity, water, co2 = values.split(",")
    assert sum(row.electricity_kwh for row in rows) == Decimal(electricity)
    assert sum(row.water_m3 for row in rows) == Decimal(water)
    assert sum(row.co2_tco2e for row in rows) == Decimal(co2)


def test_invalid_workbook_reports_every_fault_the_pack_expects():
    """D05 end to end against the error list the dataset pack ships.

    Both faults sit in different worksheets, so stopping at the first one hid the
    second until the uploader had fixed and re-uploaded. The refusal now names
    every row-level fault at once, and this asserts the sheet, date, field and
    error code against the pack's own expectations rather than against our
    wording.
    """
    base = Path(__file__).resolve().parents[3] / "test_datasets" / "workbooks"
    path = base / "D05_Bayview_2026_Q1_INVALID.xlsx"
    expected_csv = base / "expected" / "D05_expected_validation_errors.csv"
    if not path.exists() or not expected_csv.exists():
        pytest.skip(f"Workbook fixture missing: {path}")
    confirmed = MappingSuggestion(mappings=[
        ColumnMapping(field="reporting_period", worksheet="Carbon_Register", header="Month", source_unit="month"),
        ColumnMapping(field="electricity_kwh", worksheet="Electricity_Log", header="Main Incoming Grid (kWh)",
                      period_column_ref="A", source_unit="kWh", aggregation="SUM_DAILY"),
        ColumnMapping(field="water_m3", worksheet="Water_Readings", header="Water Meter - cumulative (m3)",
                      period_column_ref="A", source_unit="m3", aggregation="DIFFERENCE_CUMULATIVE",
                      opening_reading=Decimal("23410")),
        ColumnMapping(field="co2_tco2e", worksheet="Carbon_Register",
                      header="Reported Scope 2 Electricity Emissions (tCO2e)", period_column_ref="A",
                      source_unit="tCO2e", additional_column_refs=["C"]),
    ])
    with pytest.raises(ETLValidationError) as raised:
        process_confirmed_mapping(path.read_bytes(), ".xlsx", confirmed)

    rows = list(csv.DictReader(expected_csv.read_text(encoding="utf-8-sig").splitlines()))
    expected = [(row["sheet"], row["row_date"], row["field"], row["error_code"]) for row in rows]
    assert expected, "the expected-errors fixture is empty"
    assert [(issue.worksheet, issue.row_date, issue.column, issue.code)
            for issue in raised.value.issues] == expected
    # The uploader reads the message, so it must carry the whole list too.
    message = str(raised.value)
    assert f"{len(expected)} data problems" in message
    assert message.count("\n- ") == len(expected)
    assert "Electricity_Log row 20" in message and "Water_Readings row 41" in message


def test_one_decreasing_reading_is_not_blamed_on_every_later_row():
    """A single meter dip is one fault, not a cascade.

    After the dip the readings resume rising from the lower value, so judging
    each row against the one written before it reports the dip once. Comparing
    against the highest reading seen would condemn the whole rest of the month.
    """
    base = Path(__file__).resolve().parents[3] / "test_datasets" / "workbooks"
    path = base / "D05_Bayview_2026_Q1_INVALID.xlsx"
    if not path.exists():
        pytest.skip(f"Workbook fixture missing: {path}")
    confirmed = MappingSuggestion(mappings=[
        ColumnMapping(field="reporting_period", worksheet="Carbon_Register", header="Month", source_unit="month"),
        ColumnMapping(field="electricity_kwh", worksheet="Electricity_Log", header="Level 01 Submeter (kWh)",
                      period_column_ref="A", source_unit="kWh", aggregation="SUM_DAILY"),
        ColumnMapping(field="water_m3", worksheet="Water_Readings", header="Water Meter - cumulative (m3)",
                      period_column_ref="A", source_unit="m3", aggregation="DIFFERENCE_CUMULATIVE",
                      opening_reading=Decimal("23410")),
        ColumnMapping(field="co2_tco2e", worksheet="Carbon_Register",
                      header="Reported Scope 2 Electricity Emissions (tCO2e)", period_column_ref="A",
                      source_unit="tCO2e"),
    ])
    with pytest.raises(ETLValidationError) as raised:
        process_confirmed_mapping(path.read_bytes(), ".xlsx", confirmed)
    decreasing = [issue for issue in raised.value.issues if issue.code == "DECREASING_CUMULATIVE"]
    assert [(issue.row, issue.row_date) for issue in decreasing] == [(41, "2026-02-09")]


def test_cumulative_water_without_an_opening_reading_is_refused():
    path = Path(__file__).resolve().parents[3] / "test_datasets" / "workbooks" / "D01_Bayview_2026_Q1_VALID.xlsx"
    if not path.exists():
        pytest.skip(f"Workbook fixture missing: {path}")
    confirmed = MappingSuggestion(mappings=[
        ColumnMapping(field="reporting_period", worksheet="Carbon_Register", header="Month", source_unit="month"),
        ColumnMapping(field="electricity_kwh", worksheet="Electricity_Log", header="Main Incoming Grid (kWh)",
                      period_column_ref="A", source_unit="kWh", aggregation="SUM_DAILY"),
        ColumnMapping(field="water_m3", worksheet="Water_Readings", header="Water Meter - cumulative (m3)",
                      period_column_ref="A", source_unit="m3", aggregation="DIFFERENCE_CUMULATIVE"),
        ColumnMapping(field="co2_tco2e", worksheet="Carbon_Register",
                      header="Reported Scope 2 Electricity Emissions (tCO2e)", period_column_ref="A",
                      source_unit="tCO2e"),
    ])
    with pytest.raises(ETLValidationError, match="opening reading"):
        process_confirmed_mapping(path.read_bytes(), ".xlsx", confirmed)


@pytest.mark.anyio
async def test_sdk_is_called_for_unknown_structure_and_schema_is_validated():
    fake_client = MagicMock()
    fake_client.aio.models.generate_content = AsyncMock(return_value=MagicMock(text=suggestion().model_dump_json()))
    fake_client.aio.aclose = AsyncMock()
    with patch("google.genai.Client", return_value=fake_client) as client_factory:
        result = await suggest_mapping(inspect_sheets(UNKNOWN, ".csv", sample_only=True), "test-key", "test-model")
    assert result.mappings[1].header == "TNB Usage (MWh)"
    client_factory.assert_called_once()
    fake_client.aio.models.generate_content.assert_awaited_once()


@pytest.mark.anyio
@pytest.mark.parametrize("failure", [TimeoutError(), RuntimeError("rate limit")])
async def test_sdk_failure_is_bounded_and_does_not_leak_secret(failure):
    fake_client = MagicMock()
    fake_client.aio.models.generate_content = AsyncMock(side_effect=failure)
    fake_client.aio.aclose = AsyncMock()
    with patch("google.genai.Client", return_value=fake_client):
        with pytest.raises(MappingUnavailable) as raised:
            await suggest_mapping(inspect_sheets(UNKNOWN, ".csv", sample_only=True), "test-key", "test-model")
    assert "test-key" not in str(raised.value)
    assert fake_client.aio.models.generate_content.await_count == 1


@pytest.mark.anyio
async def test_ai_mapped_submission_preserves_original_and_mapping(db, storage, settings):
    org = db.query(Organisation).first()
    project = Project(name="AI Mapping Building", organisation_id=org.id)
    db.add(project)
    db.commit()
    result = await process_submission(
        upload("unfamiliar.csv", UNKNOWN, "text/csv"), "", "", "2026-Q1", uploader_id(db),
        db=db, storage=storage, settings=settings, project_id=project.id, organisation_id=org.id,
        mapping_json=suggestion().model_dump_json(),
    )
    submission = db.get(Submission, result.submission_id)
    assert result.input_format == "MONTHLY_AI_MAPPED"
    assert submission.mapping_json
    assert storage.retrieve(submission.original_storage_key) == UNKNOWN
    assert b"2026-01,30000,850,14.2" in storage.retrieve(submission.processed_storage_key)
    assert db.query(MonthlyRecord).count() == 3
