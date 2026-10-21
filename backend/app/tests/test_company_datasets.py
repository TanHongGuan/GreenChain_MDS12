"""The Sonwoy and Monosh fixtures are only useful if they still import as stated.

Each workbook ships the quarterly totals it should produce. Checking the importer
against them guards both directions: a regression in the ETL, and a regeneration
that quietly changes what the fixtures claim.
"""

import csv
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from backend.app.etl.ai_mapping import ColumnMapping, MappingSuggestion, process_confirmed_mapping
from backend.app.etl.errors import ETLValidationError
from backend.app.etl.monthly import require_submission_quarter

DATASETS = Path(__file__).resolve().parents[3] / "test_datasets"


def _opening_reading(path: Path) -> Decimal:
    """The reading before the quarter, from the file's own rules sheet.

    Read from the workbook rather than hardcoded, so a regeneration that broke
    the chain between quarters surfaces here instead of being papered over.
    """
    book = load_workbook(path, data_only=True)
    for key, value in book["Meter_And_ETL_Rules"].iter_rows(min_row=2, values_only=True):
        if key == "Water meter reading before this quarter":
            return Decimal(str(value))
    raise AssertionError(f"{path.name} does not state its opening water reading")


def _tower(path: Path) -> MappingSuggestion:
    return MappingSuggestion(mappings=[
        ColumnMapping(field="reporting_period", worksheet="Carbon_Monthly", header="Month", source_unit="month"),
        ColumnMapping(field="electricity_kwh", worksheet="Electricity_Daily", header="Landlord Supply (kWh)",
                      period_column_ref="A", source_unit="kWh", aggregation="SUM_DAILY"),
        ColumnMapping(field="water_m3", worksheet="Water_Meter", header="Bulk Meter - cumulative (m3)",
                      period_column_ref="A", source_unit="m3", aggregation="DIFFERENCE_CUMULATIVE",
                      opening_reading=_opening_reading(path)),
        ColumnMapping(field="co2_tco2e", worksheet="Carbon_Monthly",
                      header="Reported Scope 2 Emissions (tCO2e)", period_column_ref="A",
                      source_unit="tCO2e", additional_column_refs=["C"]),
    ])


def _hub(_path: Path) -> MappingSuggestion:
    return MappingSuggestion(mappings=[
        ColumnMapping(field="reporting_period", worksheet="Emissions_Register",
                      header="Reporting Month", source_unit="month"),
        ColumnMapping(field="electricity_kwh", worksheet="Electricity_Consumption",
                      header="Metered Consumption (MWh)", period_column_ref="A",
                      source_unit="MWh", aggregation="SUM_DAILY"),
        ColumnMapping(field="water_m3", worksheet="Water_Invoices", header="Water Consumed (Litres)",
                      period_column_ref="A", source_unit="L", aggregation="SUM_DAILY"),
        ColumnMapping(field="co2_tco2e", worksheet="Emissions_Register",
                      header="Reported Emissions (kgCO2e)", period_column_ref="A", source_unit="kgCO2e"),
    ])


def _science(path: Path) -> MappingSuggestion:
    return MappingSuggestion(mappings=[
        ColumnMapping(field="reporting_period", worksheet="Carbon_Register", header="Month", source_unit="month"),
        ColumnMapping(field="electricity_kwh", worksheet="Power_Log", header="Building Intake (kWh)",
                      period_column_ref="A", source_unit="kWh", aggregation="SUM_DAILY"),
        ColumnMapping(field="water_m3", worksheet="Water_Readings", header="Site Meter - cumulative (m3)",
                      period_column_ref="A", source_unit="m3", aggregation="DIFFERENCE_CUMULATIVE",
                      opening_reading=_opening_reading(path)),
        ColumnMapping(field="co2_tco2e", worksheet="Carbon_Register",
                      header="Reported Scope 2 Emissions (tCO2e)", period_column_ref="A",
                      source_unit="tCO2e", additional_column_refs=["C"]),
    ])


def _library(_path: Path) -> MappingSuggestion:
    return MappingSuggestion(mappings=[
        ColumnMapping(field="reporting_period", worksheet="Utility_Summary",
                      header="Billing Month", source_unit="month"),
        ColumnMapping(field="electricity_kwh", worksheet="Utility_Summary", header="Electricity (kWh)",
                      period_column_ref="A", source_unit="kWh"),
        ColumnMapping(field="water_m3", worksheet="Utility_Summary", header="Water (m3)",
                      period_column_ref="A", source_unit="m3"),
        ColumnMapping(field="co2_tco2e", worksheet="Utility_Summary", header="Emissions (tCO2e)",
                      period_column_ref="A", source_unit="tCO2e"),
    ])


def _riverside(_path: Path) -> MappingSuggestion:
    return MappingSuggestion(mappings=[
        ColumnMapping(field="reporting_period", worksheet="Carbon_Monthly", header="Month", source_unit="month"),
        ColumnMapping(field="electricity_kwh", worksheet="Electricity_Daily", header="Building Intake (kWh)",
                      period_column_ref="A", source_unit="kWh", aggregation="SUM_DAILY"),
        ColumnMapping(field="water_m3", worksheet="Water_Daily", header="Water Used (m3)",
                      period_column_ref="A", source_unit="m3", aggregation="SUM_DAILY"),
        ColumnMapping(field="co2_tco2e", worksheet="Carbon_Monthly",
                      header="Reported Scope 2 Emissions (tCO2e)", period_column_ref="A",
                      source_unit="tCO2e", additional_column_refs=["C"]),
    ])


def _engineering(path: Path) -> MappingSuggestion:
    return MappingSuggestion(mappings=[
        ColumnMapping(field="reporting_period", worksheet="Carbon_Register", header="Month", source_unit="month"),
        ColumnMapping(field="electricity_kwh", worksheet="Power_Log", header="Main Incoming Supply (kWh)",
                      period_column_ref="A", source_unit="kWh", aggregation="SUM_DAILY"),
        ColumnMapping(field="water_m3", worksheet="Water_Meter", header="Site Meter - cumulative (m3)",
                      period_column_ref="A", source_unit="m3", aggregation="DIFFERENCE_CUMULATIVE",
                      opening_reading=_opening_reading(path)),
        ColumnMapping(field="co2_tco2e", worksheet="Carbon_Register",
                      header="Reported Scope 2 Emissions (tCO2e)", period_column_ref="A",
                      source_unit="tCO2e", additional_column_refs=["C"]),
    ])


MAPPINGS = {
    "SONWOY_TOWER": _tower, "SONWOY_HUB": _hub, "SONWOY_RIVERSIDE": _riverside,
    "MONOSH_SCIENCE": _science, "MONOSH_LIBRARY": _library, "MONOSH_ENGINEERING": _engineering,
}

VALID = (
    [("sonwoy", f"SONWOY_TOWER_2026_Q{q}") for q in (1, 2, 3, 4)]
    + [("sonwoy", f"SONWOY_HUB_2026_Q{q}") for q in (1, 2, 3, 4)]
    + [("sonwoy", f"SONWOY_RIVERSIDE_{y}_{q}") for y, q in ((2025, "Q4"), (2026, "Q1"), (2026, "Q2"), (2026, "Q3"))]
    + [("monosh", f"MONOSH_SCIENCE_2026_Q{q}") for q in (1, 2, 3, 4)]
    + [("monosh", f"MONOSH_LIBRARY_2026_Q{q}") for q in (1, 2, 3, 4)]
    + [("monosh", f"MONOSH_ENGINEERING_{y}_{q}") for y, q in ((2025, "Q4"), (2026, "Q1"), (2026, "Q2"), (2026, "Q3"))]
)


def _builder(name: str):
    for prefix, build in MAPPINGS.items():
        if name.startswith(prefix):
            return build
    raise AssertionError(f"no mapping known for {name}")


@pytest.mark.parametrize("folder,name", VALID)
def test_dataset_imports_to_the_totals_it_ships(folder: str, name: str):
    path = DATASETS / folder / f"{name}.xlsx"
    expected_csv = DATASETS / folder / "expected" / f"{name}_expected_quarterly.csv"
    if not path.exists():
        pytest.skip(f"Dataset fixture missing: {path}")

    rows = process_confirmed_mapping(path.read_bytes(), ".xlsx", _builder(name)(path))

    parts = name.split("_")
    period = f"{parts[2]}-{parts[3]}"
    # Refuses a quarter that is short a month or counts one twice.
    require_submission_quarter(rows, period)

    expected = next(csv.DictReader(expected_csv.read_text(encoding="utf-8-sig").splitlines()))
    assert expected["period"] == period
    assert sum(row.electricity_kwh for row in rows) == Decimal(expected["electricity_kwh"])
    assert sum(row.water_m3 for row in rows) == Decimal(expected["water_m3"])
    assert sum(row.co2_tco2e for row in rows) == Decimal(expected["co2_tco2e"])


@pytest.mark.parametrize("folder,prefix,quarters", [
    ("sonwoy", "SONWOY_TOWER", ["2026_Q1", "2026_Q2", "2026_Q3", "2026_Q4"]),
    ("monosh", "MONOSH_SCIENCE", ["2026_Q1", "2026_Q2", "2026_Q3", "2026_Q4"]),
    ("monosh", "MONOSH_ENGINEERING", ["2025_Q4", "2026_Q1", "2026_Q2", "2026_Q3"]),
])
def test_a_cumulative_meter_opens_where_the_previous_quarter_closed(folder, prefix, quarters):
    """A gap between quarters is water belonging to neither.

    That is how a year's consumption silently stops adding up, so the chain is
    asserted rather than assumed - including across the calendar-year boundary.
    """
    sheet = {"SONWOY_TOWER": "Water_Meter", "MONOSH_SCIENCE": "Water_Readings",
             "MONOSH_ENGINEERING": "Water_Meter"}[prefix]
    paths = [DATASETS / folder / f"{prefix}_{quarter}.xlsx" for quarter in quarters]
    if not all(path.exists() for path in paths):
        pytest.skip("Dataset fixtures missing")

    for earlier, later in zip(paths, paths[1:]):
        book = load_workbook(earlier, data_only=True)
        closing = Decimal(str(list(book[sheet].iter_rows(values_only=True))[-1][1]))
        assert _opening_reading(later) == closing, f"{later.name} does not open where {earlier.name} closed"


def test_the_invalid_dataset_is_refused_with_every_planted_fault():
    name = "MONOSH_ENGINEERING_2026_Q3_INVALID"
    path = DATASETS / "monosh" / f"{name}.xlsx"
    expected_csv = DATASETS / "monosh" / "expected" / f"{name}_expected_errors.csv"
    if not path.exists():
        pytest.skip(f"Dataset fixture missing: {path}")

    with pytest.raises(ETLValidationError) as raised:
        process_confirmed_mapping(path.read_bytes(), ".xlsx", _engineering(path))

    rows = list(csv.DictReader(expected_csv.read_text(encoding="utf-8-sig").splitlines()))
    expected = {(row["sheet"], row["row_date"], row["field"], row["error_code"]) for row in rows}
    assert len(expected) == 3
    # All three in one refusal, not one per upload.
    assert {(issue.worksheet, issue.row_date, issue.column, issue.code)
            for issue in raised.value.issues} == expected
