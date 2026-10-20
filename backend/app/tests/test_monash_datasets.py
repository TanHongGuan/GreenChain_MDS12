"""The Monash fixtures are only useful if they still import to the stated totals.

These workbooks are generated (test_datasets/monash/generate.py) and ship with the
quarterly figures they should produce. Checking the importer against those figures
guards both directions: a regression in the ETL, and a regeneration that quietly
changes what the fixtures claim.
"""

import csv
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from backend.app.etl.ai_mapping import ColumnMapping, MappingSuggestion, process_confirmed_mapping
from backend.app.etl.errors import ETLValidationError
from backend.app.etl.monthly import require_submission_quarter

DATASETS = Path(__file__).resolve().parents[3] / "test_datasets" / "monash"
EXPECTED = DATASETS / "expected"

ENGINEERING = ["MUM_ENG_2026_Q1_APPROVED", "MUM_ENG_2026_Q1_UNREVIEWED",
               "MUM_ENG_2026_Q2_VALID", "MUM_ENG_2026_Q3_VALID", "MUM_ENG_2026_Q4_VALID"]
PHARMACY = [f"MUM_PHA_2026_Q{quarter}_VALID" for quarter in (1, 2, 3, 4)]


def _opening_reading(path: Path) -> Decimal:
    """The reading before the quarter, taken from the file's own rules sheet.

    Read from the workbook rather than hardcoded, so a regeneration that broke the
    chain between quarters would surface here instead of being papered over.
    """
    book = load_workbook(path, data_only=True)
    for key, value in book["Meter_And_ETL_Rules"].iter_rows(min_row=2, values_only=True):
        if key == "Water meter reading before this quarter":
            return Decimal(str(value))
    raise AssertionError(f"{path.name} does not state its opening water reading")


def _engineering_mapping(path: Path) -> MappingSuggestion:
    return MappingSuggestion(mappings=[
        ColumnMapping(field="reporting_period", worksheet="Carbon_Monthly", header="Month", source_unit="month"),
        ColumnMapping(field="electricity_kwh", worksheet="Electricity_Daily",
                      header="Main Incoming Supply (kWh)", period_column_ref="A",
                      source_unit="kWh", aggregation="SUM_DAILY"),
        ColumnMapping(field="water_m3", worksheet="Water_Meter", header="Bulk Meter - cumulative (m3)",
                      period_column_ref="A", source_unit="m3", aggregation="DIFFERENCE_CUMULATIVE",
                      opening_reading=_opening_reading(path)),
        # Scope 2 and Scope 1 are both reported and both belong in the total.
        ColumnMapping(field="co2_tco2e", worksheet="Carbon_Monthly",
                      header="Reported Scope 2 Emissions (tCO2e)", period_column_ref="A",
                      source_unit="tCO2e", additional_column_refs=["C"]),
    ])


def _pharmacy_mapping(_path: Path) -> MappingSuggestion:
    return MappingSuggestion(mappings=[
        ColumnMapping(field="reporting_period", worksheet="Emissions_Register",
                      header="Reporting Month", source_unit="month"),
        ColumnMapping(field="electricity_kwh", worksheet="Electricity_Consumption",
                      header="Metered Consumption (MWh)", period_column_ref="A",
                      source_unit="MWh", aggregation="SUM_DAILY"),
        ColumnMapping(field="water_m3", worksheet="Water_Invoices", header="Water Consumed (Litres)",
                      period_column_ref="A", source_unit="L", aggregation="SUM_DAILY"),
        ColumnMapping(field="co2_tco2e", worksheet="Emissions_Register",
                      header="Reported Emissions (kgCO2e)", period_column_ref="A", source_unit="kgCO2e"),
    ])


@pytest.mark.parametrize("name", ENGINEERING + PHARMACY)
def test_dataset_imports_to_the_totals_it_ships(name: str):
    path = DATASETS / f"{name}.xlsx"
    expected_csv = EXPECTED / f"{name}_expected_quarterly.csv"
    if not path.exists():
        pytest.skip(f"Dataset fixture missing: {path}")

    build = _engineering_mapping if "_ENG_" in name else _pharmacy_mapping
    rows = process_confirmed_mapping(path.read_bytes(), ".xlsx", build(path))

    quarter = name.split("_")[3]
    period = f"2026-{quarter}"
    # Refuses a quarter that is short a month or counts one twice.
    require_submission_quarter(rows, period)

    expected = next(csv.DictReader(expected_csv.read_text(encoding="utf-8-sig").splitlines()))
    assert expected["period"] == period
    assert sum(row.electricity_kwh for row in rows) == Decimal(expected["electricity_kwh"])
    assert sum(row.water_m3 for row in rows) == Decimal(expected["water_m3"])
    assert sum(row.co2_tco2e for row in rows) == Decimal(expected["co2_tco2e"])


def test_the_engineering_water_meter_runs_continuously_across_the_year():
    """Each quarter must open on the reading the previous one closed with.

    A gap between quarters is water that belongs to neither, which is how a year's
    consumption silently stops adding up. The shipped pack this replaces had
    exactly that break between its two quarters.
    """
    quarters = ["MUM_ENG_2026_Q1_APPROVED", "MUM_ENG_2026_Q2_VALID",
                "MUM_ENG_2026_Q3_VALID", "MUM_ENG_2026_Q4_VALID"]
    if not all((DATASETS / f"{name}.xlsx").exists() for name in quarters):
        pytest.skip("Dataset fixtures missing")

    for earlier, later in zip(quarters, quarters[1:]):
        book = load_workbook(DATASETS / f"{earlier}.xlsx", data_only=True)
        closing = Decimal(str(list(book["Water_Meter"].iter_rows(values_only=True))[-1][1]))
        assert _opening_reading(DATASETS / f"{later}.xlsx") == closing, (
            f"{later} does not open where {earlier} closed"
        )


def test_the_invalid_dataset_is_refused_with_every_planted_fault():
    name = "MUM_ENG_2026_Q1_INVALID"
    path = DATASETS / f"{name}.xlsx"
    expected_csv = EXPECTED / f"{name}_expected_errors.csv"
    if not path.exists():
        pytest.skip(f"Dataset fixture missing: {path}")

    with pytest.raises(ETLValidationError) as raised:
        process_confirmed_mapping(path.read_bytes(), ".xlsx", _engineering_mapping(path))

    rows = list(csv.DictReader(expected_csv.read_text(encoding="utf-8-sig").splitlines()))
    expected = {(row["sheet"], row["row_date"], row["field"], row["error_code"]) for row in rows}
    assert len(expected) == 3
    # All three in one refusal, not one per upload.
    assert {(issue.worksheet, issue.row_date, issue.column, issue.code)
            for issue in raised.value.issues} == expected
