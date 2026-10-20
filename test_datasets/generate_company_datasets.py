"""Generate the Sonwoy and Monosh company test workbooks.

SYNTHETIC DATA. Sonwoy and Monosh are invented organisations. Nothing in these
files is a real meter reading, bill or emissions statement for any company,
and every workbook says so on its own 00_Read_Me sheet.

Run:  python test_datasets/generate_company_datasets.py

Two organisations, two buildings each, four quarters of 2026 per building -
sixteen workbooks. The four buildings are written in four different house
styles, because a real platform receives whatever the producer wrote:

                            electricity              water                 carbon
  Sonwoy Corporate Tower    daily kWh  DD/MM/YYYY    cumulative m3 ISO      monthly tCO2e
  Sonwoy Logistics Hub      daily MWh  MM/DD/YYYY    daily litres YYYY/MM/DD  monthly kgCO2e
  Monosh Science Building   daily kWh  YYYY-MM-DD    cumulative m3 DD/MM/YYYY monthly tCO2e
  Monosh Library            monthly kWh              monthly m3             monthly tCO2e

Between them they cover every unit conversion the importer performs, all three
aggregation rules and every date order it has to resolve. No column mixes units,
so each file imports as written.
"""

from __future__ import annotations

import csv
import random
from calendar import monthrange
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parent
YEAR = 2026
QUARTERS = {"Q1": (1, 2, 3), "Q2": (4, 5, 6), "Q3": (7, 8, 9), "Q4": (10, 11, 12)}

CLASSIFICATION = "SYNTHETIC - generated test data, NOT real meter readings or statements"

# Occupancy drives consumption. A commercial tower is steadier than a campus; the
# warehouse barely notices the calendar at all.
OFFICE_LOAD = {1: 0.92, 2: 0.95, 3: 1.02, 4: 1.04, 5: 1.06, 6: 1.05,
               7: 1.07, 8: 1.08, 9: 1.04, 10: 1.02, 11: 0.98, 12: 0.84}
CAMPUS_LOAD = {1: 0.74, 2: 0.98, 3: 1.06, 4: 1.08, 5: 1.05, 6: 0.86,
               7: 0.92, 8: 1.04, 9: 1.07, 10: 1.09, 11: 0.99, 12: 0.70}
FLAT_LOAD = {month: 1.0 for month in range(1, 13)}

# Peninsular Malaysia grid emission factor, tCO2e per MWh, as published by the
# Energy Commission (Suruhanjaya Tenaga). Carbon is still reported per month in
# these files rather than derived by the importer.
GRID_FACTOR = Decimal("0.585")


def days_in(months: tuple[int, ...]) -> list[date]:
    return [date(YEAR, month, day)
            for month in months
            for day in range(1, monthrange(YEAR, month)[1] + 1)]


def daily_series(rng: random.Random, days: list[date], base: float,
                 load: dict[int, float], weekend: float) -> list[Decimal]:
    return [Decimal(f"{base * load[day.month] * (weekend if day.weekday() >= 5 else 1.0) * rng.uniform(0.94, 1.06):.2f}")
            for day in days]


def read_me(book: Workbook, rows: list[tuple[str, str]]) -> None:
    sheet = book.create_sheet("00_Read_Me")
    sheet.append(["Key", "Value"])
    for key, value in rows:
        sheet.append([key, value])


def rules(book: Workbook, rows: list[tuple[str, str]]) -> None:
    # Keys are prose, not codes. The importer's note extraction drops
    # identifier-shaped tokens ("GRID-A") before the sheet reaches the model,
    # which would leave each rule without its subject.
    sheet = book.create_sheet("Meter_And_ETL_Rules")
    sheet.append(["Source field / meter", "Interpretation / action"])
    for key, value in rows:
        sheet.append([key, value])


def standard_read_me(book: Workbook, *, organisation: str, building: str, quarter: str,
                     notes: str) -> None:
    read_me(book, [
        ("Dataset", f"{building} {YEAR}-{quarter}"),
        ("Organisation", organisation),
        ("Building", building),
        ("Reporting period", f"{YEAR}-{quarter}"),
        ("Scenario", "VALID"),
        ("Data classification", CLASSIFICATION),
        ("Scope", "Building-level operational electricity, water and reported CO2-equivalent emissions"),
        ("Notes", notes),
    ])


def corporate_tower(quarter: str, rng: random.Random, opening: Decimal,
                    organisation: str, building: str) -> tuple[Workbook, dict, Decimal]:
    """Daily kWh day-first, a cumulative water meter in ISO dates, monthly tCO2e."""
    months = QUARTERS[quarter]
    days = days_in(months)
    electricity = daily_series(rng, days, 4_180.0, OFFICE_LOAD, 0.46)
    water_daily = daily_series(rng, days, 18.6, OFFICE_LOAD, 0.46)

    book = Workbook()
    power = book.active
    power.title = "Electricity_Daily"
    power.append(["Reading Date (DD/MM/YYYY)", "Landlord Supply (kWh)", "Tenant Submeter (kWh)",
                  "Meter Point", "Notes"])
    for day, value in zip(days, electricity):
        power.append([day.strftime("%d/%m/%Y"), str(value),
                      str((value * Decimal("0.38")).quantize(Decimal("0.01"))),
                      "Landlord intake", "Landlord supply INCLUDES the tenant submeter"])

    water = book.create_sheet("Water_Meter")
    water.append(["Reading Date", "Bulk Meter - cumulative (m3)", "Meter Serial", "Remarks"])
    running = opening
    for day, value in zip(days, water_daily):
        running += value
        water.append([day.strftime("%Y-%m-%d"), str(running), "Tower bulk meter",
                      "Cumulative reading, not daily usage"])

    carbon = book.create_sheet("Carbon_Monthly")
    carbon.append(["Month", "Reported Scope 2 Emissions (tCO2e)", "Reported Scope 1 Emissions (tCO2e)",
                   "Reporting Basis", "Statement Ref"])
    totals_co2 = Decimal("0")
    for month in months:
        month_kwh = sum(value for day, value in zip(days, electricity) if day.month == month)
        two = ((month_kwh / 1000) * GRID_FACTOR).quantize(Decimal("0.001"))
        one = (two * Decimal("0.035")).quantize(Decimal("0.001"))
        totals_co2 += two + one
        carbon.append([f"{YEAR}-{month:02d}", str(two), str(one),
                       "Reported totals, not derived by GreenChain", f"Statement {YEAR}-{month:02d}"])

    standard_read_me(book, organisation=organisation, building=building, quarter=quarter,
                     notes="Expected to pass all data checks")
    rules(book, [
        ("Landlord Supply (kWh)", "Daily consumption on the building intake; sum the days for a month"),
        ("Tenant Submeter (kWh)", "Already included in the landlord supply; DO NOT ADD"),
        ("Bulk Meter - cumulative (m3)", "Cumulative reading; a month's usage is the rise across it"),
        ("Water meter reading before this quarter", str(opening)),
        ("Reported Scope 2 Emissions (tCO2e)", "Add to Scope 1 for the building total; do not derive from electricity"),
        ("Date parsing", "Electricity DD/MM/YYYY, water YYYY-MM-DD, carbon YYYY-MM"),
    ])

    return book, {"electricity_kwh": sum(electricity), "water_m3": running - opening,
                  "co2_tco2e": totals_co2}, running


def logistics_hub(quarter: str, rng: random.Random, organisation: str, building: str) -> tuple[Workbook, dict]:
    """Daily MWh month-first, daily litres in slashed ISO, monthly kgCO2e."""
    months = QUARTERS[quarter]
    days = days_in(months)
    electricity_mwh = [Decimal(f"{float(value) / 1000:.4f}")
                       for value in daily_series(rng, days, 6_240.0, FLAT_LOAD, 0.82)]
    water_litres = [Decimal(f"{float(value):.0f}")
                    for value in daily_series(rng, days, 9_400.0, FLAT_LOAD, 0.82)]

    book = Workbook()
    power = book.active
    power.title = "Electricity_Consumption"
    power.append(["Usage Date (MM/DD/YYYY)", "Metered Consumption (MWh)", "Tariff Code", "Invoice Reference"])
    for day, value in zip(days, electricity_mwh):
        power.append([day.strftime("%m/%d/%Y"), str(value), "E2 Medium Voltage",
                      f"Invoice {YEAR}-{day.month:02d}"])

    water = book.create_sheet("Water_Invoices")
    water.append(["Date Recorded", "Water Consumed (Litres)", "Service Area", "Bill Number"])
    for day, value in zip(days, water_litres):
        water.append([day.strftime("%Y/%m/%d"), str(value), "Whole site", f"Water bill {YEAR}-{day.month:02d}"])

    carbon = book.create_sheet("Emissions_Register")
    carbon.append(["Reporting Month", "Reported Emissions (kgCO2e)", "Basis", "Register Ref"])
    total_kg = Decimal("0")
    for month in months:
        month_mwh = sum(value for day, value in zip(days, electricity_mwh) if day.month == month)
        kilograms = (month_mwh * GRID_FACTOR * 1000 * Decimal("1.04")).quantize(Decimal("0.01"))
        total_kg += kilograms
        carbon.append([f"{YEAR}-{month:02d}", str(kilograms), "Reported operational total",
                       f"Register {YEAR}-{month:02d}"])

    standard_read_me(book, organisation=organisation, building=building, quarter=quarter,
                     notes="Expected to pass all data checks")
    rules(book, [
        ("Metered Consumption (MWh)", "Daily consumption in megawatt-hours; multiply by 1,000 for kWh"),
        ("Water Consumed (Litres)", "Daily volume in litres; divide by 1,000 for m3"),
        ("Reported Emissions (kgCO2e)", "Reported monthly total in kilograms; divide by 1,000 for tCO2e"),
        ("Date parsing", "Electricity MM/DD/YYYY, water YYYY/MM/DD, emissions YYYY-MM"),
    ])

    return book, {"electricity_kwh": sum(electricity_mwh) * 1000,
                  "water_m3": sum(water_litres) / 1000,
                  "co2_tco2e": total_kg / 1000}


def science_building(quarter: str, rng: random.Random, opening: Decimal,
                     organisation: str, building: str) -> tuple[Workbook, dict, Decimal]:
    """Daily kWh in ISO dates, a cumulative water meter written day-first, monthly tCO2e."""
    months = QUARTERS[quarter]
    days = days_in(months)
    electricity = daily_series(rng, days, 7_950.0, CAMPUS_LOAD, 0.61)
    water_daily = daily_series(rng, days, 29.8, CAMPUS_LOAD, 0.61)

    book = Workbook()
    power = book.active
    power.title = "Power_Log"
    power.append(["Date", "Building Intake (kWh)", "Laboratory Submeter (kWh)", "Feed", "Comment"])
    for day, value in zip(days, electricity):
        power.append([day.strftime("%Y-%m-%d"), str(value),
                      str((value * Decimal("0.44")).quantize(Decimal("0.01"))),
                      "Main intake", "Intake INCLUDES the laboratory submeter"])

    water = book.create_sheet("Water_Readings")
    water.append(["Reading Date (DD/MM/YYYY)", "Site Meter - cumulative (m3)", "Meter Serial", "Remarks"])
    running = opening
    for day, value in zip(days, water_daily):
        running += value
        water.append([day.strftime("%d/%m/%Y"), str(running), "Site bulk meter",
                      "Cumulative reading, not daily usage"])

    carbon = book.create_sheet("Carbon_Register")
    carbon.append(["Month", "Reported Scope 2 Emissions (tCO2e)", "Reported Other Emissions (tCO2e)",
                   "Reporting Basis", "Source Ref"])
    totals_co2 = Decimal("0")
    for month in months:
        month_kwh = sum(value for day, value in zip(days, electricity) if day.month == month)
        two = ((month_kwh / 1000) * GRID_FACTOR).quantize(Decimal("0.001"))
        other = (two * Decimal("0.061")).quantize(Decimal("0.001"))
        totals_co2 += two + other
        carbon.append([f"{YEAR}-{month:02d}", str(two), str(other),
                       "Reported totals, not derived by GreenChain", f"Source {YEAR}-{month:02d}"])

    standard_read_me(book, organisation=organisation, building=building, quarter=quarter,
                     notes="Expected to pass all data checks")
    rules(book, [
        ("Building Intake (kWh)", "Daily consumption on the main intake; sum the days for a month"),
        ("Laboratory Submeter (kWh)", "Already included in the intake; DO NOT ADD"),
        ("Site Meter - cumulative (m3)", "Cumulative reading; a month's usage is the rise across it"),
        ("Water meter reading before this quarter", str(opening)),
        ("Reported Scope 2 Emissions (tCO2e)", "Add to Other emissions for the building total"),
        ("Date parsing", "Electricity YYYY-MM-DD, water DD/MM/YYYY, carbon YYYY-MM"),
    ])

    return book, {"electricity_kwh": sum(electricity), "water_m3": running - opening,
                  "co2_tco2e": totals_co2}, running


def monthly_bills(quarter: str, rng: random.Random, organisation: str, building: str) -> tuple[Workbook, dict]:
    """One row per month for all three measures - a building that only has its bills."""
    months = QUARTERS[quarter]
    book = Workbook()
    sheet = book.active
    sheet.title = "Utility_Summary"
    sheet.append(["Billing Month", "Electricity (kWh)", "Water (m3)", "Emissions (tCO2e)",
                  "Account Reference"])
    totals = {"electricity_kwh": Decimal("0"), "water_m3": Decimal("0"), "co2_tco2e": Decimal("0")}
    for month in months:
        days = monthrange(YEAR, month)[1]
        kwh = Decimal(f"{2_140.0 * days * CAMPUS_LOAD[month] * rng.uniform(0.96, 1.04):.2f}")
        cubic = Decimal(f"{9.4 * days * CAMPUS_LOAD[month] * rng.uniform(0.96, 1.04):.2f}")
        carbon = ((kwh / 1000) * GRID_FACTOR).quantize(Decimal("0.001"))
        totals["electricity_kwh"] += kwh
        totals["water_m3"] += cubic
        totals["co2_tco2e"] += carbon
        sheet.append([f"{YEAR}-{month:02d}", str(kwh), str(cubic), str(carbon),
                      f"Account {YEAR}-{month:02d}"])

    standard_read_me(book, organisation=organisation, building=building, quarter=quarter,
                     notes="Already monthly: one row per month, nothing to aggregate")
    rules(book, [
        ("Electricity (kWh)", "Already a monthly total; do not sum days"),
        ("Water (m3)", "Already a monthly total in cubic metres"),
        ("Emissions (tCO2e)", "Reported monthly total in tonnes"),
        ("Date parsing", "All columns YYYY-MM"),
    ])
    return book, totals


def write_expected(folder: Path, name: str, building: str, period: str, totals: dict) -> None:
    (folder / "expected").mkdir(parents=True, exist_ok=True)
    with (folder / "expected" / f"{name}_expected_quarterly.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["building", "period", "electricity_kwh", "water_m3", "co2_tco2e"])
        writer.writerow([building, period, totals["electricity_kwh"], totals["water_m3"], totals["co2_tco2e"]])


def build_cumulative(folder: Path, prefix: str, organisation: str, building: str,
                     maker, opening: Decimal, seed: str) -> None:
    """A building whose water meter runs continuously through the year.

    Each quarter opens on the reading the one before it closed with, so the four
    files chain instead of leaving water that belongs to no quarter.
    """
    running = opening
    for quarter in QUARTERS:
        rng = random.Random(f"{seed}-{quarter}")
        book, totals, running = maker(quarter, rng, running, organisation, building)
        name = f"{prefix}_{YEAR}_{quarter}"
        book.save(folder / f"{name}.xlsx")
        write_expected(folder, name, building, f"{YEAR}-{quarter}", totals)


def build_simple(folder: Path, prefix: str, organisation: str, building: str, maker, seed: str) -> None:
    for quarter in QUARTERS:
        rng = random.Random(f"{seed}-{quarter}")
        book, totals = maker(quarter, rng, organisation, building)
        name = f"{prefix}_{YEAR}_{quarter}"
        book.save(folder / f"{name}.xlsx")
        write_expected(folder, name, building, f"{YEAR}-{quarter}", totals)


def main() -> None:
    sonwoy = ROOT / "sonwoy"
    monosh = ROOT / "monosh"
    for folder in (sonwoy, monosh):
        folder.mkdir(parents=True, exist_ok=True)

    build_cumulative(sonwoy, "SONWOY_TOWER", "Sonwoy", "Sonwoy Corporate Tower",
                     corporate_tower, Decimal("18260.40"), "sonwoy-tower")
    build_simple(sonwoy, "SONWOY_HUB", "Sonwoy", "Sonwoy Logistics Hub",
                 logistics_hub, "sonwoy-hub")
    build_cumulative(monosh, "MONOSH_SCIENCE", "Monosh", "Monosh Science Building",
                     science_building, Decimal("57410.80"), "monosh-science")
    build_simple(monosh, "MONOSH_LIBRARY", "Monosh", "Monosh Library",
                 monthly_bills, "monosh-library")

    for folder in (sonwoy, monosh):
        print(f"{folder.name}: {len(list(folder.glob('*.xlsx')))} workbooks")


if __name__ == "__main__":
    main()
