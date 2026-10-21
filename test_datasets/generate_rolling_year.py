"""Generate a rolling reporting year, 2025-Q4 through 2026-Q3.

SYNTHETIC DATA. Sonwoy and Monosh are invented organisations and these are not
real meter readings, bills or emissions statements for any company. Every figure
comes from this script and a fixed seed, and each workbook says so on its own
00_Read_Me sheet.

Run:  python test_datasets/generate_rolling_year.py

One new building per organisation, four consecutive quarters each, so the set
spans a year that crosses the calendar boundary:

  Sonwoy Riverside Office   daily kWh DD/MM/YYYY, daily m3, monthly tCO2e
  Monosh Engineering Hall   daily kWh YYYY-MM-DD, cumulative m3, monthly tCO2e

Plus one deliberately broken quarter that must be refused.

New buildings rather than new quarters of the existing ones, so nothing here
collides with the files already in test_datasets/sonwoy and
test_datasets/monosh. Shared shaping (teaching and office load curves, the grid
emission factor, the read-me and rules sheets) is imported rather than copied.
"""

from __future__ import annotations

import csv
import random
from calendar import monthrange
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook

from generate_company_datasets import (
    CAMPUS_LOAD, CLASSIFICATION, GRID_FACTOR, OFFICE_LOAD, QUARTERS,
    daily_series, read_me, rules, write_expected,
)

ROOT = Path(__file__).resolve().parent

# A reporting year that crosses the calendar boundary, which is how a financial
# year usually sits against a calendar one.
PERIODS = ((2025, "Q4"), (2026, "Q1"), (2026, "Q2"), (2026, "Q3"))


def days_in(year: int, months: tuple[int, ...]) -> list[date]:
    return [date(year, month, day)
            for month in months
            for day in range(1, monthrange(year, month)[1] + 1)]


def _read_me(book: Workbook, *, organisation: str, building: str, year: int, quarter: str,
             scenario: str, notes: str) -> None:
    read_me(book, [
        ("Dataset", f"{building} {year}-{quarter}"),
        ("Organisation", organisation),
        ("Building", building),
        ("Reporting period", f"{year}-{quarter}"),
        ("Scenario", scenario),
        ("Data classification", CLASSIFICATION),
        ("Scope", "Building-level operational electricity, water and reported CO2-equivalent emissions"),
        ("Notes", notes),
    ])


def riverside_office(year: int, quarter: str, rng: random.Random,
                     organisation: str, building: str) -> tuple[Workbook, dict]:
    """A let office: daily kWh written day-first, daily water, monthly carbon.

    About 9,000 m2 at roughly 145 kWh/m2/year, which is ordinary for a Malaysian
    commercial building, and a weekend that drops to just over a third of a
    working day because only base services run.
    """
    months = QUARTERS[quarter]
    days = days_in(year, months)
    electricity = daily_series(rng, days, 3_575.0, OFFICE_LOAD, 0.38)
    water = daily_series(rng, days, 14.2, OFFICE_LOAD, 0.34)

    book = Workbook()
    power = book.active
    power.title = "Electricity_Daily"
    power.append(["Reading Date (DD/MM/YYYY)", "Building Intake (kWh)", "Car Park Submeter (kWh)",
                  "Meter Point", "Notes"])
    for day, value in zip(days, electricity):
        power.append([day.strftime("%d/%m/%Y"), str(value),
                      str((value * Decimal("0.11")).quantize(Decimal("0.01"))),
                      "Main intake", "Intake INCLUDES the car park submeter"])

    tap = book.create_sheet("Water_Daily")
    tap.append(["Reading Date (DD/MM/YYYY)", "Water Used (m3)", "Zone", "Remarks"])
    for day, value in zip(days, water):
        tap.append([day.strftime("%d/%m/%Y"), str(value), "Whole building",
                    "Daily consumption, not a meter reading"])

    carbon = book.create_sheet("Carbon_Monthly")
    carbon.append(["Month", "Reported Scope 2 Emissions (tCO2e)", "Reported Scope 1 Emissions (tCO2e)",
                   "Reporting Basis", "Statement Ref"])
    total_co2 = Decimal("0")
    for month in months:
        month_kwh = sum(value for day, value in zip(days, electricity) if day.month == month)
        scope2 = ((month_kwh / 1000) * GRID_FACTOR).quantize(Decimal("0.001"))
        scope1 = (scope2 * Decimal("0.028")).quantize(Decimal("0.001"))
        total_co2 += scope2 + scope1
        carbon.append([f"{year}-{month:02d}", str(scope2), str(scope1),
                       "Reported totals, not derived by GreenChain", f"Statement {year}-{month:02d}"])

    _read_me(book, organisation=organisation, building=building, year=year, quarter=quarter,
             scenario="VALID", notes="Expected to pass all data checks")
    rules(book, [
        ("Building Intake (kWh)", "Daily consumption on the main intake; sum the days for a month"),
        ("Car Park Submeter (kWh)", "Already included in the intake; DO NOT ADD"),
        ("Water Used (m3)", "Daily consumption in cubic metres; sum the days for a month"),
        ("Reported Scope 2 Emissions (tCO2e)", "Add to Scope 1 for the building total"),
        ("Date parsing", "Electricity DD/MM/YYYY, water DD/MM/YYYY, carbon YYYY-MM"),
    ])

    return book, {"electricity_kwh": sum(electricity), "water_m3": sum(water), "co2_tco2e": total_co2}


def engineering_hall(year: int, quarter: str, rng: random.Random, opening: Decimal,
                     organisation: str, building: str) -> tuple[Workbook, dict, Decimal]:
    """A teaching and laboratory building: ISO dates and a cumulative water meter.

    About 14,000 m2 at roughly 175 kWh/m2/year, with the teaching calendar
    driving the shape - December and January are quiet, the exam months are not.
    """
    months = QUARTERS[quarter]
    days = days_in(year, months)
    electricity = daily_series(rng, days, 6_700.0, CAMPUS_LOAD, 0.57)
    water_daily = daily_series(rng, days, 27.4, CAMPUS_LOAD, 0.57)

    book = Workbook()
    power = book.active
    power.title = "Power_Log"
    power.append(["Date", "Main Incoming Supply (kWh)", "Workshop Submeter (kWh)", "Feed", "Comment"])
    for day, value in zip(days, electricity):
        power.append([day.strftime("%Y-%m-%d"), str(value),
                      str((value * Decimal("0.36")).quantize(Decimal("0.01"))),
                      "Main intake", "Main supply INCLUDES the workshop submeter"])

    water = book.create_sheet("Water_Meter")
    water.append(["Reading Date", "Site Meter - cumulative (m3)", "Meter Serial", "Remarks"])
    running = opening
    for day, value in zip(days, water_daily):
        running += value
        water.append([day.strftime("%Y-%m-%d"), str(running), "Site bulk meter",
                      "Cumulative reading, not daily usage"])

    carbon = book.create_sheet("Carbon_Register")
    carbon.append(["Month", "Reported Scope 2 Emissions (tCO2e)", "Reported Other Emissions (tCO2e)",
                   "Reporting Basis", "Source Ref"])
    total_co2 = Decimal("0")
    for month in months:
        month_kwh = sum(value for day, value in zip(days, electricity) if day.month == month)
        scope2 = ((month_kwh / 1000) * GRID_FACTOR).quantize(Decimal("0.001"))
        other = (scope2 * Decimal("0.055")).quantize(Decimal("0.001"))
        total_co2 += scope2 + other
        carbon.append([f"{year}-{month:02d}", str(scope2), str(other),
                       "Reported totals, not derived by GreenChain", f"Source {year}-{month:02d}"])

    _read_me(book, organisation=organisation, building=building, year=year, quarter=quarter,
             scenario="VALID", notes="Expected to pass all data checks")
    rules(book, [
        ("Main Incoming Supply (kWh)", "Daily consumption on the main intake; sum the days for a month"),
        ("Workshop Submeter (kWh)", "Already included in the main supply; DO NOT ADD"),
        ("Site Meter - cumulative (m3)", "Cumulative reading; a month's usage is the rise across it"),
        ("Water meter reading before this quarter", str(opening)),
        ("Reported Scope 2 Emissions (tCO2e)", "Add to Other emissions for the building total"),
        ("Date parsing", "Electricity YYYY-MM-DD, water YYYY-MM-DD, carbon YYYY-MM"),
    ])

    return book, {"electricity_kwh": sum(electricity), "water_m3": running - opening,
                  "co2_tco2e": total_co2}, running


def break_it(book: Workbook) -> list[tuple[str, str, str, str]]:
    """Plant three faults a correct importer must find, and say where they are.

    Cells are damaged, never whole rows: a blank row is skipped as padding, so
    blanking one would remove the fault rather than create it.
    """
    faults = []

    power = book["Power_Log"]
    row = 17
    power.cell(row=row, column=2, value="NO READING")
    faults.append(("Power_Log", str(power.cell(row=row, column=1).value),
                   "Main Incoming Supply (kWh)", "INVALID_NUMERIC"))

    row = 55
    # Assigning through .value: Worksheet.cell(value=None) leaves the cell alone,
    # so passing None there silently fails to create the blank.
    power.cell(row=row, column=2).value = None
    faults.append(("Power_Log", str(power.cell(row=row, column=1).value),
                   "Main Incoming Supply (kWh)", "MISSING_VALUE"))

    water = book["Water_Meter"]
    row = 34
    previous = Decimal(str(water.cell(row=row - 1, column=2).value))
    water.cell(row=row, column=2, value=str(previous - Decimal("112.80")))
    faults.append(("Water_Meter", str(water.cell(row=row, column=1).value),
                   "Site Meter - cumulative (m3)", "DECREASING_CUMULATIVE"))

    return faults


def main() -> None:
    sonwoy = ROOT / "sonwoy"
    monosh = ROOT / "monosh"

    for year, quarter in PERIODS:
        rng = random.Random(f"riverside-{year}-{quarter}")
        book, totals = riverside_office(year, quarter, rng, "Sonwoy", "Sonwoy Riverside Office")
        name = f"SONWOY_RIVERSIDE_{year}_{quarter}"
        book.save(sonwoy / f"{name}.xlsx")
        write_expected(sonwoy, name, "Sonwoy Riverside Office", f"{year}-{quarter}", totals)

    # The water meter runs continuously through the reporting year, so each
    # quarter opens on the reading the one before it closed with.
    running = Decimal("63480.20")
    for year, quarter in PERIODS:
        rng = random.Random(f"engineering-{year}-{quarter}")
        book, totals, running = engineering_hall(year, quarter, rng, running, "Monosh",
                                                 "Monosh Engineering Hall")
        name = f"MONOSH_ENGINEERING_{year}_{quarter}"
        book.save(monosh / f"{name}.xlsx")
        write_expected(monosh, name, "Monosh Engineering Hall", f"{year}-{quarter}", totals)

    # One quarter regenerated and deliberately damaged. Built from its own seed so
    # it is not mistaken for a corrupted copy of the valid Q3 beside it.
    broken_rng = random.Random("engineering-invalid")
    broken, _, _ = engineering_hall(2026, "Q3", broken_rng, Decimal("63480.20"), "Monosh",
                                    "Monosh Engineering Hall")
    _read_me(broken, organisation="Monosh", building="Monosh Engineering Hall", year=2026,
             quarter="Q3", scenario="INVALID",
             notes="Contains deliberate faults; must be refused, not published.")
    faults = break_it(broken)
    broken.save(monosh / "MONOSH_ENGINEERING_2026_Q3_INVALID.xlsx")
    with (monosh / "expected" / "MONOSH_ENGINEERING_2026_Q3_INVALID_expected_errors.csv").open(
            "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["severity", "sheet", "row_date", "field", "error_code", "expected_handling"])
        for sheet, row_date, field, code in faults:
            writer.writerow(["ERROR", sheet, row_date, field, code, "Refuse the upload; publish nothing"])

    print(f"sonwoy: {len(list(sonwoy.glob('SONWOY_RIVERSIDE_*.xlsx')))} rolling-year workbooks")
    print(f"monosh: {len(list(monosh.glob('MONOSH_ENGINEERING_*.xlsx')))} rolling-year workbooks "
          f"(one of them invalid)")


if __name__ == "__main__":
    main()
